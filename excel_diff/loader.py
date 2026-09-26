"""Excel dosyalarını okuma."""
from __future__ import annotations

import io

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .model import Table
from .normalize import norm_ws, to_text


def _open(data: bytes):
    return load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def sheet_names(data: bytes) -> list[str]:
    wb = _open(data)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _raw_rows(data: bytes, sheet: str | None) -> list[list[str]]:
    wb = _open(data)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return [[to_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def detect_header_row(rows: list[list[str]]) -> int:
    """İlk en az 2 dolu hücresi olan satırı başlık kabul eder (1 tabanlı)."""
    for i, row in enumerate(rows[:50]):
        if sum(1 for v in row if norm_ws(v)) >= 2:
            return i + 1
    return 1


def read_table(data: bytes, name: str, sheet: str | None = None, header_row: int | None = None) -> Table:
    rows = _raw_rows(data, sheet)
    if not rows:
        return Table(name, sheet or "", [], [], [])
    if header_row is None:
        header_row = detect_header_row(rows)
    header_idx = max(0, header_row - 1)
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header = rows[header_idx]
    body = rows[header_idx + 1:]
    excel_rows = list(range(header_idx + 2, header_idx + 2 + len(body)))

    # Tamamen boş sütunları at
    keep = [c for c in range(width) if norm_ws(header[c]) or any(norm_ws(r[c]) for r in body)]
    columns: list[str] = []
    seen: dict[str, int] = {}
    for c in keep:
        col = norm_ws(header[c]) or f"Sütun {get_column_letter(c + 1)}"
        if col in seen:
            seen[col] += 1
            col = f"{col} ({seen[col]})"
        else:
            seen[col] = 1
        columns.append(col)

    out_rows, out_excel = [], []
    for r, xr in zip(body, excel_rows):
        vals = [r[c] for c in keep]
        if any(norm_ws(v) for v in vals):  # tamamen boş satırları at
            out_rows.append(vals)
            out_excel.append(xr)
    return Table(name, sheet or "", columns, out_rows, out_excel)
